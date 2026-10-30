"""N3 scoped participation review and read-only handoff regressions."""

import csv
import hashlib
import gc
import io
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from provider_directory_truth_layer.intake_api import create_workbench
from provider_directory_truth_layer.intake_store import IntakeStore
from provider_directory_truth_layer.operations_sql import run as run_operations


SOURCE = "payer:n3-synthetic"
SCOPE = "fixture-participation"
FIELDS = ["record_id", "npi", "provider_name", "address", "city", "state", "zip", "phone",
          "location_id", "network_code", "organization_code", "participation_status",
          "effective_date", "end_date"]


@pytest.fixture
def store(tmp_path):
    return IntakeStore(tmp_path / "intake.sqlite3", tmp_path / "runtime")


def row(network="CEDAR", location="NORTH", record_id="LIST-1", **changes):
    return {"record_id": record_id, "npi": "1000000001", "provider_name": "Synthetic Provider",
            "address": "10 Fiction St", "city": "Example", "state": "TN", "zip": "01234",
            "phone": "5550100000", "location_id": location, "network_code": network,
            "organization_code": "CLINIC", "participation_status": "listed",
            "effective_date": "2026-01-01", "end_date": "", **changes}


def data(rows):
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=FIELDS)
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode()


def setup(store):
    payer = store.create_catalog("organizations", {"label": "Synthetic Payer", "kind": "payer", "actor": "op", "reason": "Fixture"})
    clinic = store.create_catalog("organizations", {"label": "Synthetic Clinic", "kind": "clinic", "actor": "op", "reason": "Fixture"})
    catalogs = {
        "CEDAR": store.create_catalog("networks", {"label": "CEDAR", "payer_organization_id": payer["id"], "actor": "op", "reason": "Fixture"}),
        "PINE": store.create_catalog("networks", {"label": "PINE", "payer_organization_id": payer["id"], "actor": "op", "reason": "Fixture"}),
        "NORTH": store.create_catalog("locations", {"label": "NORTH", "organization_id": clinic["id"], "actor": "op", "reason": "Fixture"}),
        "SOUTH": store.create_catalog("locations", {"label": "SOUTH", "organization_id": clinic["id"], "actor": "op", "reason": "Fixture"}),
        "CLINIC": clinic,
    }
    bindings = {"network": {}, "location": {}, "organization": {}}
    for kind, codes in (("network", ["CEDAR", "PINE"]), ("location", ["NORTH", "SOUTH"]), ("organization", ["CLINIC"])):
        for code in codes:
            bindings[kind][code] = store.bind_source({"source": SOURCE, "kind": kind, "source_key": code,
                                                      "target_id": catalogs[code]["id"], "actor": "op", "reason": "Fixture"})["id"]
    return catalogs, bindings


def receive(store, rows, *, completeness="bounded", scope=SCOPE):
    return store.receive(data(rows), "synthetic-network.csv", source=SOURCE,
                         metadata={"scope": scope, "completeness": completeness, "as_of": "2026-02-01"})


def mapped(store, receipt, bindings):
    return store.mapping(receipt["id"], {"mapping": {field: field for field in FIELDS},
                                         "roster_kind": "network", "bindings": bindings})


def accept(store, receipt):
    return store.commit(receipt["id"], {"actor": "op", "reason": "Reviewed source assertion",
                                        "policy": "valid_rows", "expected_attempt": receipt["attempt_id"]})


def initial_snapshot(store, bindings):
    first = receive(store, [row(), row("CEDAR", "SOUTH"), row("PINE", "NORTH")])
    return accept(store, mapped(store, first, bindings))


def _hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _sheet_rows(book, name):
    return list(book[name].values)[1:]


def test_missing_blocked_unresolved_and_changed_review_preserve_human_work(store):
    catalogs, bindings = setup(store)
    first = initial_snapshot(store, bindings)
    cedar_case = next(case for case in first["linked_cases"] if
                      json.loads(case["evidence"])["normalized"]["network_code"] == "CEDAR"
                      and json.loads(case["evidence"])["normalized"]["location_id"] == "NORTH")
    original = store.case(cedar_case["id"])
    assigned = store.event(cedar_case["id"], {"action": "assign", "owner": "reviewer", "actor": "op",
                                               "reason": "Contact clinic", "expected_version": 0})
    incoming_rows = [row(phone="5550100008"),
                     row("CEDAR", "SOUTH", effective_date="bad-date"),
                     row("UNKNOWN", "NORTH", record_id="UNKNOWN")]
    incoming = receive(store, incoming_rows)
    incoming = mapped(store, incoming, bindings)
    comparison = incoming["preview"]["comparison"]
    assert comparison["changed"] == comparison["blocked_invalid"] == comparison["unresolved"] == 1
    assert comparison["missing_from_this_delivery"] == 1
    assert incoming["preview"]["valid_count"] == 1
    assert incoming["preview"]["quarantined_count"] == 2
    accepted = accept(store, incoming)
    case = store.case(cedar_case["id"])
    assert case["status"] == assigned["status"]
    assert case["owner"] == "reviewer"
    assert case["version"] == assigned["version"]
    assert case["history"] == assigned["history"]
    assert case["evidence"] == original["evidence"]
    assert len(case["later_evidence"]) == 1
    assert case["later_evidence"][0]["human_verified_at"] is None
    work = store.work()["items"]
    assert len(work) == 5  # Three initial identities, one known invalid, one missing; no unknown-network case.
    validation = [item for item in work if item["field_name"] == "participation_validation"]
    missing = [item for item in work if item["field_name"] == "participation_missing"]
    assert len(validation) == len(missing) == 1
    assert validation[0]["network_id"] == catalogs["CEDAR"]["id"]
    assert missing[0]["network_id"] == catalogs["PINE"]["id"]
    assert all(item["status"] == "new" for item in validation + missing)
    for item in validation + missing:
        detail = store.case(item["id"])
        assert detail["evidence"]["observation_id"] is None
        assert detail["participation_review"]
        assert detail["scope"]["network_id"] == item["network_id"]
        assert detail["evidence"]["receipt_id"] == accepted["id"]
        assert detail["evidence"]["baseline_receipt_id"] == first["id"]
        review = detail["participation_review"][0]
        assert review["receipt_id"] == accepted["id"]
        assert review["payload"]["comparison"]["previous_receipt_id"] == first["id"]
        lineage = run_operations(store.path, "lineage", case_id=item["id"])
        assert len(lineage) == 1
        assert lineage[0]["observation_id"] is None
        assert lineage[0]["receipt_id"] == accepted["id"]
        assert lineage[0]["comparison_status"] in {"missing", "blocked_invalid"}
    validation_detail = store.case(validation[0]["id"])
    assert validation_detail["evidence"]["original"]["effective_date"] == "bad-date"
    assert validation_detail["participation_review"][0]["payload"]["source_row"]["locator"] == "row:3"
    missing_detail = store.case(missing[0]["id"])
    assert missing_detail["participation_review"][0]["payload"]["source_row"] is None
    assert missing_detail["participation_review"][0]["payload"]["baseline_observation"]["id"] == missing_detail["evidence"]["baseline_observation_id"]
    assert all(item["field_name"] != "participation" or item["network_id"] in
               {catalogs["CEDAR"]["id"], catalogs["PINE"]["id"]} for item in work)
    assert any(item["id"] == validation[0]["id"] for item in accepted["linked_cases"])
    assert any(item["id"] == missing[0]["id"] for item in accepted["linked_cases"])
    assert (store.artifacts / accepted["artifact_sha"]).read_bytes() == data(incoming_rows)
    client = TestClient(create_workbench(store.path, store.root))
    for item in validation + missing:
        response = client.get("/api/work-export.json", params={"case_id": item["id"]})
        assert response.status_code == 200
        packet = response.json()
        assert [case["id"] for case in packet["cases"]] == [item["id"]]
        assert packet["participations"] == []  # Quarantine/absence is review evidence, not accepted membership.
        assert packet["evidence"] == []
        assert len(packet["review_evidence"]) == 1
        assert packet["review_evidence"][0]["receipt_id"] == accepted["id"]
        assert {receipt["id"] for receipt in packet["receipts"]} >= {first["id"], accepted["id"]}
        assert all(network["id"] == item["network_id"] for network in packet["catalog"]["networks"])


def test_scoped_work_and_exports_exclude_unrelated_network_rows(store):
    catalogs, bindings = setup(store)
    first = initial_snapshot(store, bindings)
    cedar_id, pine_id = catalogs["CEDAR"]["id"], catalogs["PINE"]["id"]
    cedar = store.work(network_id=cedar_id)["items"]
    pine = store.work(network_id=pine_id)["items"]
    assert len(cedar) == 2 and len(pine) == 1
    assert all(item["network_id"] == cedar_id for item in cedar)
    assert all(item["network_id"] == pine_id for item in pine)
    assert len(store.work(location_id=catalogs["NORTH"]["id"])["items"]) == 2
    assert len(store.work(organization_id=catalogs["CLINIC"]["id"])["items"]) == 3
    assert store.work(network_id=cedar_id, location_id=catalogs["SOUTH"]["id"])["items"][0]["location_label"] == "SOUTH"
    client = TestClient(create_workbench(store.path, store.root))
    artifact = store.artifacts / first["artifact_sha"]
    before = (_hash(store.path), _hash(artifact))
    gc.collect()  # Prior request connections must already be closed; export cannot checkpoint a late writer.
    response = client.get("/api/work-export.json", params={"network_id": cedar_id})
    assert response.status_code == 200
    snapshot = response.json()
    assert snapshot["network_id"] == cedar_id
    assert {"cases", "evidence", "participations", "review_evidence", "catalog", "bindings",
            "receipts", "attempts", "decisions", "relationships"} <= set(snapshot)
    assert {item["id"] for item in snapshot["cases"]} == {item["id"] for item in cedar}
    assert all(item["network_id"] == cedar_id for item in snapshot["participations"])
    assert all(item["id"] != pine_id for item in snapshot["catalog"]["networks"])
    assert all((item.get("target_id") or item.get("network_id")) == cedar_id
               for item in snapshot["bindings"] if item.get("kind") == "network")
    assert any(item["id"] == first["id"] for item in snapshot["receipts"])
    assert any(item["id"] == first["attempt_id"] for item in snapshot["attempts"])
    response = client.get("/api/work-export.xlsx", params={"network_id": cedar_id})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content), read_only=True, data_only=False)
    assert {"Read me", "Cases", "Evidence", "Receipts", "Attempts", "Decisions", "Relationships", "Payloads",
            "Participations", "Review evidence", "Catalog", "Bindings"} <= set(book.sheetnames)
    assert len(_sheet_rows(book, "Cases")) == len(cedar)
    assert len(_sheet_rows(book, "Participations")) == 2
    assert len(_sheet_rows(book, "Review evidence")) >= 2
    assert all(not (item[0] == "networks" and item[1] == pine_id) for item in _sheet_rows(book, "Catalog"))
    assert all(item[4] == cedar_id for item in _sheet_rows(book, "Bindings") if item[2] == "network")
    book.close()
    selected = cedar[0]["id"]
    case_snapshot = client.get("/api/work-export.json", params={"case_id": selected})
    assert case_snapshot.status_code == 200
    body = case_snapshot.json()
    assert {item["id"] for item in body["cases"]} == {selected}
    assert all(item["case_id"] == selected for item in body["evidence"])
    assert all(item["case_id"] == selected for item in body["participations"])
    assert all(item["case_id"] == selected for item in body["review_evidence"])
    assert all(item["case_id"] == selected for item in body["decisions"])
    assert all(item["case_id"] == selected for item in body["relationships"])
    case_book_response = client.get("/api/work-export.xlsx", params={"case_id": selected})
    assert case_book_response.status_code == 200
    case_book = load_workbook(io.BytesIO(case_book_response.content), read_only=True, data_only=False)
    assert len(_sheet_rows(case_book, "Cases")) == 1
    assert len(_sheet_rows(case_book, "Participations")) == 1
    case_book.close()
    assert (_hash(store.path), _hash(artifact)) == before
    all_work = client.get("/api/work-export.json")
    assert all_work.status_code == 200
    assert len(all_work.json()["cases"]) == len(store.work()["items"])
    assert client.get("/api/work-export.json", params={"case_id": "unknown"}).status_code == 404
    assert client.get("/api/work-export.json", params={"network_id": "unknown"}).status_code == 404
    assert client.get("/api/work-export.xlsx", params={"case_id": cedar[0]["id"], "network_id": cedar_id}).status_code == 422


def test_repeats_do_not_duplicate_review_evidence_and_restart_preserves_it(store):
    _, bindings = setup(store)
    first = initial_snapshot(store, bindings)
    incoming_rows = [row(phone="5550100008"), row("CEDAR", "SOUTH", effective_date="bad-date")]
    second = accept(store, mapped(store, receive(store, incoming_rows), bindings))
    with store.db() as conn:
        before = conn.execute("SELECT count(*) FROM participation_review_evidence").fetchone()[0]
    repeat = accept(store, mapped(store, receive(store, incoming_rows), bindings))
    assert repeat["state"] == "duplicate_delivery"
    with store.db() as conn:
        assert conn.execute("SELECT count(*) FROM participation_review_evidence").fetchone()[0] == before
    restarted = IntakeStore(store.path, store.root)
    assert {item["id"] for item in restarted.work()["items"]} == {item["id"] for item in store.work()["items"]}
    assert (restarted.artifacts / first["artifact_sha"]).read_bytes() == data([row(), row("CEDAR", "SOUTH"), row("PINE", "NORTH")])
    assert (restarted.artifacts / second["artifact_sha"]).read_bytes() == data(incoming_rows)


def test_explicit_full_empty_snapshot_records_absence_only(store):
    _, bindings = setup(store)
    first = initial_snapshot(store, bindings)
    empty = receive(store, [], completeness="full")
    empty = mapped(store, empty, bindings)
    assert empty["preview"]["comparison"]["baseline"]["attempt_id"] == first["attempt_id"]
    assert empty["preview"]["comparison"]["missing_from_this_delivery"] == 3
    accepted = accept(store, empty)
    assert accepted["state"] == "committed"
    assert accepted["observations"] == []
    assert all(case["field_name"] == "participation_missing" for case in accepted["linked_cases"])
    assert len(accepted["linked_cases"]) == 3
    assert all(case["status"] == "new" for case in accepted["linked_cases"])


def test_all_invalid_bounded_network_delivery_cannot_create_review_work(store):
    _, bindings = setup(store)
    first = initial_snapshot(store, bindings)
    before = {item["id"] for item in store.work()["items"]}
    incoming = receive(store, [row(effective_date="bad-date")])
    incoming = mapped(store, incoming, bindings)
    assert incoming["preview"]["valid_count"] == 0
    assert incoming["preview"]["comparison"]["blocked_invalid"] == 1
    with pytest.raises(ValueError, match="No valid rows"):
        accept(store, incoming)
    assert {item["id"] for item in store.work()["items"]} == before
    assert first["observations"] == store.detail(first["id"])["observations"]


def test_sql_compare_uses_full_network_key_and_rejects_incompatible_rosters(store):
    _, bindings = setup(store)
    first = initial_snapshot(store, bindings)
    later = receive(store, [row(phone="5550100008"), row("CEDAR", "SOUTH"), row("PINE", "NORTH")])
    later = accept(store, mapped(store, later, bindings))
    differences = run_operations(store.path, "compare", left_attempt=first["attempt_id"],
                                 right_attempt=later["attempt_id"])
    assert len(differences) == 3
    assert sorted(item["difference"] for item in differences) == ["changed", "unchanged", "unchanged"]
    assert len({item["left_observation_id"] for item in differences}) == 3
    clinic = receive(store, [row()])
    clinic = store.mapping(clinic["id"], {"mapping": {field: field for field in FIELDS}})
    clinic = accept(store, clinic)
    with pytest.raises(ValueError, match="roster|scope|compatible"):
        run_operations(store.path, "compare", left_attempt=first["attempt_id"],
                       right_attempt=clinic["attempt_id"])
    cedar_only = {**bindings, "network": {"CEDAR": bindings["network"]["CEDAR"]}}
    subset = receive(store, [row(phone="5550100009"), row("CEDAR", "SOUTH")])
    subset = accept(store, mapped(store, subset, cedar_only))
    with pytest.raises(ValueError, match="network|scope|compatible"):
        run_operations(store.path, "compare", left_attempt=first["attempt_id"],
                       right_attempt=subset["attempt_id"])


def test_identical_bytes_other_dataset_do_not_reuse_mapping_baseline_or_export_context(store):
    catalogs, bindings = setup(store)
    source_rows = [row(), row("PINE", "NORTH")]
    first = receive(store, source_rows)
    first = mapped(store, first, bindings)
    revision = store.save_mapping(first["id"], {"name": "Reviewed network roster", "actor": "op",
                                               "reason": "Fixture", "expected_attempt": first["attempt_id"]})
    assert revision["id"] in {item["id"] for item in store.saved_mappings(first["id"])["items"]}
    first = accept(store, first)
    other = receive(store, source_rows, scope="unrelated-dataset")
    assert other["artifact_sha"] == first["artifact_sha"]
    assert other["duplicate_of"] is None
    assert store.saved_mappings(other["id"])["items"] == []
    assert other["attempts"][0]["spec"].get("roster_kind", "clinic") == "clinic"
    assert "bindings" not in other["attempts"][0]["spec"]
    assert other["preview"]["comparison"]["baseline"] is None
    other = mapped(store, other, bindings)
    assert other["preview"]["comparison"]["baseline"] is None
    client = TestClient(create_workbench(store.path, store.root))
    selected = client.get("/api/work-export.json", params={"network_id": catalogs["CEDAR"]["id"]})
    assert selected.status_code == 200
    receipt_ids = {receipt["id"] for receipt in selected.json()["receipts"]}
    assert first["id"] in receipt_ids
    assert other["id"] not in receipt_ids


def test_selected_case_retains_catalog_and_bindings_for_later_source_organization(store):
    catalogs, bindings=setup(store)
    initial=initial_snapshot(store,bindings)
    selected=next(case for case in initial['linked_cases'] if json.loads(case['evidence'])['normalized'].get('network_code')=='CEDAR' and json.loads(case['evidence'])['normalized'].get('location_id')=='NORTH')
    original=store.case(selected['id'])
    other=store.create_catalog('organizations',{'label':'Another synthetic clinic','kind':'clinic','actor':'op','reason':'Explicit separate organization'})
    replacement=store.bind_source({'source':SOURCE,'kind':'organization','source_key':'CLINIC','target_id':other['id'],
        'expected_binding_id':bindings['organization']['CLINIC'],'actor':'op','reason':'Source now asserts another organization'})
    updated={kind:dict(values) for kind,values in bindings.items()}
    updated['organization']['CLINIC']=replacement['id']
    incoming=receive(store,[row(phone='5550100008'),row('CEDAR','SOUTH'),row('PINE','NORTH')])
    accept(store,mapped(store,incoming,updated))
    current=store.case(selected['id'])
    assert current['evidence']==original['evidence'] and current['scope']==original['scope']
    assert current['history']==original['history'] and current['status']==original['status']
    client=TestClient(create_workbench(store.path,store.root))
    exported=client.get('/api/work-export.json',params={'case_id':selected['id']}).json()
    referenced={item['organization_id'] for item in exported['participations']}
    assert referenced=={catalogs['CLINIC']['id'],other['id']}
    assert referenced <= {item['id'] for item in exported['catalog']['organizations']}
    assert referenced <= {item['target_id'] for item in exported['bindings'] if item['kind']=='organization'}
