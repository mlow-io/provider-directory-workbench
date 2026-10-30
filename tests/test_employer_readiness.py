"""Employer demonstration and handoff contracts, using only test-owned storage."""
import hashlib
import io
import json
import sqlite3
from collections import defaultdict

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from provider_directory_truth_layer import work_export
from provider_directory_truth_layer.demo_samples import SAMPLES, sample_payloads
from provider_directory_truth_layer.intake_api import create_workbench
from provider_directory_truth_layer.intake_store import IntakeStore


def provider(**changes):
    return {'record_id': '0007', 'npi': '0000000001', 'provider_name': 'Synthetic Example',
            'address': '1 Example Street', 'city': 'Boston', 'state': 'MA', 'zip': '02108',
            'phone': '6175550100', 'location_id': 'NORTH', 'reference_phone': '6175550199',
            'source_updated_at': '2026-06-26', **changes}


def accept(store, rows, source='synthetic-test', **metadata):
    receipt = store.receive(json.dumps(rows).encode(), 'synthetic.json', source=source, metadata=metadata)
    assert receipt['state'] == 'preview_ready', receipt
    return store.commit(receipt['id'], {'actor': 'Synthetic reviewer', 'reason': 'Explicitly reviewed test evidence',
        'policy': 'valid_rows', 'expected_attempt': receipt['attempt_id']})


def rows(workbook, sheet):
    data = workbook[sheet].iter_rows(values_only=True)
    headers = next(data)
    return [dict(zip(headers, row)) for row in data]


def payloads(workbook):
    parts = defaultdict(list)
    for row in rows(workbook, 'Payloads'):
        parts[row['payload_id']].append((int(row['part']), row['json_chunk']))
    return {key: json.loads(''.join(part for _, part in sorted(values))) for key, values in parts.items()}


def state(store):
    with sqlite3.connect(store.path) as conn:
        dump = '\n'.join(conn.iterdump())
    artifacts = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in store.artifacts.iterdir()}
    return dump, artifacts


def test_samples_are_opt_in_and_regular_uploads_remain_usable(tmp_path, monkeypatch):
    monkeypatch.delenv('PDTL_DEMO', raising=False)
    client = TestClient(create_workbench(tmp_path/'test.db', tmp_path))
    assert client.get('/api/demo/samples').json() == {'items': []}
    assert client.post('/api/demo/samples/clinic-workbook').status_code == 404
    ordinary = client.post('/api/intakes/files?filename=local.json&designation=local',
        content=json.dumps([provider()]).encode())
    assert ordinary.status_code == 200
    assert ordinary.json()['designation'] == 'local'
    assert ordinary.json()['state'] == 'preview_ready'
    empty = load_workbook(io.BytesIO(client.get('/api/work-export.xlsx').content))
    assert rows(empty, 'Cases') == []
    assert client.get('/api/work-export.xlsx?case_id=missing').status_code == 404


@pytest.mark.parametrize('sample_id', [sample['id'] for sample in SAMPLES])
def test_loading_any_sample_uses_retained_normal_intake_without_acceptance(tmp_path, sample_id):
    client = TestClient(create_workbench(tmp_path/'test.db', tmp_path, demo=True))
    assert {s['id'] for s in client.get('/api/demo/samples').json()['items']} == {s['id'] for s in SAMPLES}
    response = client.post(f'/api/demo/samples/{sample_id}')
    assert response.status_code == 200, response.text
    receipt = response.json()
    assert receipt['source'] == ('synthetic-network-demo' if sample_id.startswith('network-') else 'synthetic-demo')
    if sample_id.startswith('network-'):
        assert receipt['state']=='needs_mapping'
        assert client.get('/api/catalog').json()['networks']==[]
    assert receipt['designation'] == 'synthetic'
    assert receipt['metadata']['demo_sample'] == sample_id
    assert 'SYNTHETIC' in receipt['metadata']['notice']
    assert receipt['state'] in {'needs_mapping', 'preview_ready'}
    assert receipt['observations_count'] == 0
    assert client.get('/api/work').json()['items'] == []
    assert client.get(f"/api/intakes/{receipt['id']}/artifact").content == sample_payloads()[sample_id]
    if sample_id == 'clinic-workbook':
        assert receipt['state'] == 'needs_mapping'
        assert receipt['attempts'][-1]['spec']['header_row'] == 2
        assert receipt['preview']['mapping'] == {}
    assert client.post('/api/demo/samples/not-an-allowed-sample').status_code == 404


def test_guided_mapping_repeat_and_restart_preserve_owner_evidence_and_decisions(tmp_path):
    path = tmp_path/'test.db'
    client = TestClient(create_workbench(path, tmp_path, demo=True))
    receipt = client.post('/api/demo/samples/clinic-workbook').json()
    common = {'actor': 'Synthetic reviewer', 'reason': 'Review two sites, preserve malformed row', 'policy': 'valid_rows'}
    assert client.post(f"/api/intakes/{receipt['id']}/commit", json={**common,
        'expected_attempt': receipt['attempt_id']}).status_code == 422
    mapped = client.post(f"/api/intakes/{receipt['id']}/mapping", json={
        'header_row': 2, 'sheet': 'Clinic roster', 'mapping': receipt['metadata']['suggested_mapping']}).json()
    assert (mapped['preview']['valid_count'], mapped['preview']['quarantined_count']) == (2, 1)
    committed = client.post(f"/api/intakes/{receipt['id']}/commit", json={**common,
        'expected_attempt': mapped['attempt_id']}).json()
    assert committed['state'] == 'committed'
    cases = client.get('/api/work').json()['items']
    assert {c['location_id'] for c in cases} == {'synthetic-demo:NORTH', 'synthetic-demo:SOUTH'}
    north = next(c for c in cases if c['location_id'].endswith(':NORTH'))
    client.post(f"/api/work/{north['id']}/events", json={'action': 'assign', 'owner': 'Demo owner',
        'follow_up': '2026-11-01', 'actor': 'Reviewer', 'reason': 'Retain through repeat and restart', 'expected_version': 0})
    # Clear generation cache and create another service against the SAME test-owned root.
    before_bytes = sample_payloads()['clinic-workbook']
    sample_payloads.cache_clear()
    assert sample_payloads()['clinic-workbook'] == before_bytes
    resumed = TestClient(create_workbench(path, tmp_path, demo=True))
    repeated = resumed.post('/api/demo/samples/clinic-workbook').json()
    assert repeated['artifact_sha'] == committed['artifact_sha']
    repeat_commit = resumed.post(f"/api/intakes/{repeated['id']}/commit", json={**common,
        'expected_attempt': repeated['attempt_id']}).json()
    assert repeat_commit['state'] == 'duplicate_delivery'
    assert len(resumed.get('/api/work').json()['items']) == 2
    detail = resumed.get(f"/api/work/{north['id']}").json()
    assert (detail['owner'], detail['follow_up'], detail['version']) == ('Demo owner', '2026-11-01', 1)
    assert len(detail['history']) == 1
    assert detail['evidence']['locator'] == 'Clinic roster!3'
    assert detail['evidence']['normalized']['zip'] == '02108'
    packet = load_workbook(io.BytesIO(resumed.get('/api/work-export.xlsx').content))
    assert {r['receipt_id'] for r in rows(packet, 'Receipts')} == {receipt['id'], repeated['id']}


def test_xlsx_literal_identifiers_formula_nulls_and_complete_long_payloads(tmp_path):
    store = IntakeStore(tmp_path/'test.db', tmp_path)
    long_value = 'Evidence 😀\u0001' * 5000
    expanded_controls = '\u0001' * 6000 + '\ufffe'
    original = provider(provider_name='=HYPERLINK("https://example.invalid","unsafe")',
        location_id='+001', city='NA', blank=None, empty='', literal='NA', long_note=long_value,
        formula_values=['=1+1', '+1+1', '-1+1', '@SUM(A1)'])
    accepted = accept(store, [original], note=long_value)
    case = store.case(accepted['linked_cases'][0]['id'])
    case = store.event(case['id'], {'action': 'triage', 'actor': '@Reviewer', 'reason': '=1+1', 'expected_version': case['version']})
    case = store.event(case['id'], {'action': 'approve', 'actor': '+Reviewer', 'reason': long_value,
        'proposal': '6175550199', 'expected_version': case['version']})
    case = store.event(case['id'], {'action': 'submit', 'actor': 'Reviewer', 'reason': expanded_controls,
        'submission_reference': '=HYPERLINK("bad")', 'expected_version': case['version']})
    case = store.event(case['id'], {'action': 'verify', 'actor': 'Reviewer', 'reason': 'Observe synthetic mismatch',
        'observed': '6175550177', 'verification_note': long_value, 'expected_version': case['version']})
    client = TestClient(create_workbench(store.path, store.root, demo=False))
    before = state(store)
    response = client.get('/api/work-export.xlsx', params={'case_id': case['id']})
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment' in response.headers['content-disposition']
    assert state(store) == before
    workbook = load_workbook(io.BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ['Read me','Cases','Evidence','Receipts','Attempts','Decisions','Relationships','Payloads']
    exported = rows(workbook, 'Cases')[0]
    assert (exported['npi'], exported['native_record_id']) == ('0000000001', '0007')
    assert exported['provider_name'] == original['provider_name']
    assert exported['owner'] is None
    assert (exported['listed_value'], exported['requested_value'], exported['approved_value'], exported['observed_value']) == (
        '6175550100', '6175550199', '6175550199', '6175550177')
    assert exported['status'] == 'reopened'
    evidence = rows(workbook, 'Evidence')[0]
    assert (evidence['zip'], evidence['city'], evidence['source_updated_at'], evidence['human_verified_at']) == ('02108', 'NA', '2026-06-26', None)
    decoded = payloads(workbook)
    observation = decoded[evidence['observation_payload_id']]
    assert observation['original'] == original
    assert observation['original']['blank'] is None
    assert observation['original']['empty'] == ''
    decisions = rows(workbook, 'Decisions')
    assert decisions[1]['reason'].startswith('payload:')
    assert decoded[decisions[1]['reason'].removeprefix('payload:')] == long_value
    assert decisions[2]['reason'].startswith('payload:')
    assert decoded[decisions[2]['reason'].removeprefix('payload:')] == expanded_controls
    assert decoded[decisions[-1]['event_payload_id']]['payload']['verification_note'] == long_value
    assert len(rows(workbook, 'Payloads')) > len(decoded), 'Large original values must be chunked, never truncated'
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    assert cell.data_type == 's'
                    assert cell.number_format == '@'
                    assert len(cell.value.encode('utf-16-le')) // 2 <= 32767
                    assert cell.hyperlink is None


def test_export_is_complete_beyond_ui_bounds_and_selected_case_remains_scoped(tmp_path):
    store = IntakeStore(tmp_path/'test.db', tmp_path)
    accepted = accept(store, [provider(record_id=f'{index:05d}', location_id=f'SITE-{index:03d}') for index in range(205)])
    selected = accepted['linked_cases'][0]['id']
    initial = store.case(selected)
    # Fifty-five later observations share stable provider/location/native IDs.
    record = initial['evidence']['normalized']
    accept(store, [{**record, 'phone': f'617555{index:04d}'} for index in range(55)])
    for index in range(55):
        current = store.case(selected)
        store.event(selected, {'action': 'assign', 'owner': f'Owner-{index}', 'actor': 'Synthetic reviewer',
            'reason': f'Accountable handoff {index}', 'expected_version': current['version']})
    before = state(store)
    all_work = load_workbook(io.BytesIO(work_export.handoff_workbook(store.path)))
    assert len(rows(all_work, 'Cases')) == 205
    assert len(rows(all_work, 'Evidence')) == 260
    assert len(rows(all_work, 'Decisions')) == 55
    scoped = load_workbook(io.BytesIO(work_export.handoff_workbook(store.path, selected)))
    assert len(rows(scoped, 'Cases')) == 1
    evidence = rows(scoped, 'Evidence')
    assert len(evidence) == 56
    assert {r['case_id'] for r in evidence} == {selected}
    assert {r['location_id'] for r in evidence} == {initial['location_id']}
    assert len([r for r in evidence if r['role'] == 'initial']) == 1
    assert len([r for r in evidence if r['role'] == 'later']) == 55
    assert {r['row_locator'] for r in evidence}.issuperset({'$[0]', '$[54]'})
    assert len(rows(scoped, 'Decisions')) == 55
    assert {r['case_id'] for r in rows(scoped, 'Decisions')} == {selected}
    assert state(store) == before


def test_read_snapshot_cannot_mix_case_version_with_concurrent_new_decision(tmp_path, monkeypatch):
    store = IntakeStore(tmp_path/'test.db', tmp_path)
    accepted = accept(store, [provider()])
    case_id = accepted['linked_cases'][0]['id']
    called = []

    def concurrent_writer():
        called.append(True)
        store.event(case_id, {'action': 'assign', 'owner': 'Concurrent owner', 'actor': 'Synthetic writer',
            'reason': 'Write after export case SELECT', 'expected_version': 0})
        return '2026-10-26T12:00:00+00:00'

    monkeypatch.setattr(work_export, 'now', concurrent_writer)
    workbook = load_workbook(io.BytesIO(work_export.handoff_workbook(store.path, case_id)))
    assert called == [True]
    assert store.case(case_id)['version'] == 1
    assert rows(workbook, 'Cases')[0]['version'] == '0'
    assert rows(workbook, 'Cases')[0]['owner'] is None
    assert rows(workbook, 'Decisions') == []


def test_explicit_native_location_relationship_and_its_audit_survive_export(tmp_path):
    store = IntakeStore(tmp_path/'test.db', tmp_path)
    received = store.receive(sample_payloads()['nppes'], 'synthetic-nppes.json', source='synthetic-reference', profile='nppes')
    accepted = store.commit(received['id'], {'actor': 'Reviewer', 'reason': 'Retain synthetic native evidence',
        'policy': 'valid_rows', 'expected_attempt': received['attempt_id']})
    observation = accepted['observations'][0]
    store.relationship({'observation_id': observation['id'], 'provider_id': 'NPI:1234567890',
        'location_id': 'explicit:SOUTH', 'location_index': 1, 'actor': 'Relationship reviewer',
        'reason': 'Selected second native practice address explicitly'})
    selected = next(case for case in store.work()['items'] if case['location_id'] == 'explicit:SOUTH')
    workbook = load_workbook(io.BytesIO(work_export.handoff_workbook(store.path, selected['id'])))
    assert rows(workbook, 'Cases')[0]['location_id'] == 'explicit:SOUTH'
    evidence = rows(workbook, 'Evidence')[0]
    assert evidence['location_id'] == 'explicit:SOUTH'
    assert evidence['address'] == '2 Example Street'
    assert evidence['row_locator'].endswith('addresses[1]')
    assert evidence['native_observation_location_id'] is None
    relation = rows(workbook, 'Relationships')[0]
    assert relation['observation_id'] == observation['id']
    assert (relation['provider_id'], relation['location_id'], relation['location_index']) == ('NPI:1234567890', 'explicit:SOUTH', '1')
    assert relation['actor'] == 'Relationship reviewer'
    assert relation['reason'] == 'Selected second native practice address explicitly'
    decoded = payloads(workbook)
    exported = decoded[rows(workbook, 'Cases')[0]['case_payload_id']]
    assert exported['evidence']['normalized']['address'] == '2 Example Street'
