"""Regression coverage for case-scoped native location evidence in XLSX handoffs."""

import hashlib
import io
import json
import sqlite3
from collections import defaultdict

from openpyxl import load_workbook
from fastapi.testclient import TestClient

from provider_directory_truth_layer.demo_samples import sample_payloads
from provider_directory_truth_layer.intake_store import IntakeStore
from provider_directory_truth_layer.intake_api import create_workbench


def sheet_rows(workbook, name):
    values = workbook[name].iter_rows(values_only=True)
    headers = next(values)
    return [dict(zip(headers, row)) for row in values]


def decoded_payloads(workbook):
    chunks = defaultdict(list)
    for row in sheet_rows(workbook, 'Payloads'):
        chunks[row['payload_id']].append((int(row['part']), row['json_chunk']))
    return {key: json.loads(''.join(chunk for _, chunk in sorted(parts)))
            for key, parts in chunks.items()}


def durable_state(store):
    with sqlite3.connect(store.path) as conn:
        database = '\n'.join(conn.iterdump())
    artifacts = {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                 for path in store.artifacts.iterdir()}
    return database, artifacts


def test_explicit_native_location_scope_survives_later_evidence_and_export(tmp_path):
    store = IntakeStore(tmp_path / 'review.db', tmp_path)
    source = json.loads(sample_payloads()['nppes'])
    observations = []
    expected_south_phones = ('6175550101', '6175550198')

    for phone in expected_south_phones:
        delivery = json.loads(json.dumps(source))
        delivery['results'][0]['addresses'][1]['telephone_number'] = phone
        receipt = store.receive(json.dumps(delivery).encode(), 'synthetic-nppes.json',
                                source='synthetic-reference', profile='nppes')
        accepted = store.commit(receipt['id'], {
            'actor': 'Synthetic reviewer', 'reason': 'Accepted native observation',
            'policy': 'valid_rows', 'expected_attempt': receipt['attempt_id'],
        })
        observation = accepted['observations'][0]
        observations.append(observation['id'])
        for index, location in ((0, 'explicit:NORTH'), (1, 'explicit:SOUTH')):
            store.relationship({
                'observation_id': observation['id'], 'provider_id': 'NPI:1234567890',
                'location_id': location, 'location_index': index,
                'actor': 'Synthetic reviewer', 'reason': 'Selected this practice address',
            })

    cases = {case['location_id']: case for case in store.work()['items']
             if case['location_id'] in {'explicit:NORTH', 'explicit:SOUTH'}}
    assert set(cases) == {'explicit:NORTH', 'explicit:SOUTH'}
    client = TestClient(create_workbench(store.path, tmp_path, demo=False))
    before = durable_state(store)

    response = client.get('/api/work-export.xlsx', params={'case_id': cases['explicit:SOUTH']['id']})
    assert response.status_code == 200
    south = load_workbook(io.BytesIO(response.content))
    evidence = sheet_rows(south, 'Evidence')
    assert len(evidence) == 2
    assert [row['role'] for row in evidence] == ['initial', 'later']
    assert [row['observation_id'] for row in evidence] == observations
    assert [row['listed_phone'] for row in evidence] == list(expected_south_phones)
    for row in evidence:
        assert row['provider_id'] == 'NPI:1234567890'
        assert row['location_id'] == 'explicit:SOUTH'
        assert row['row_locator'] == '$.results[0].addresses[1]'
        assert (row['address'], row['zip']) == ('2 Example Street', '02109')
        assert row['native_observation_location_id'] is None
        assert row['native_observation_locator'] == '$.results[0]'

    relationships = sheet_rows(south, 'Relationships')
    assert len(relationships) == 2
    assert {(row['case_id'], row['observation_id'], row['location_id'], row['location_index'])
            for row in relationships} == {
                (cases['explicit:SOUTH']['id'], observation_id, 'explicit:SOUTH', '1')
                for observation_id in observations
            }

    payloads = decoded_payloads(south)
    for row in evidence:
        native = payloads[row['observation_payload_id']]
        assert native['id'] == row['observation_id']
        assert native['locator'] == '$.results[0]'
        assert native['location_id'] is None
        assert native['normalized']['address'] == '1 Example Street'

    response = client.get('/api/work-export.xlsx')
    assert response.status_code == 200
    all_work = load_workbook(io.BytesIO(response.content))
    scoped_relationships = sheet_rows(all_work, 'Relationships')
    assert len(scoped_relationships) == 4
    assert {(row['case_id'], row['location_id'], row['observation_id'])
            for row in scoped_relationships} == {
                (cases[location]['id'], location, observation_id)
                for location in cases for observation_id in observations
            }
    assert durable_state(store) == before
