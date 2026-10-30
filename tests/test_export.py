from datetime import datetime
from pathlib import Path
import sqlite3

from openpyxl import load_workbook
import pandas as pd

from provider_directory_truth_layer.export import REMEDIATION_OUTPUTS, REMEDIATION_TABLES
from provider_directory_truth_layer.pipeline import run_pipeline


def test_pipeline_writes_review_workbook_with_expected_sheets(tmp_path):
    repo = Path(__file__).resolve().parents[1]
    run_pipeline(
        input_csv=repo / "data/raw/sample_provider_directory.csv",
        config_path=repo / "config/qa_rules.yml",
        output_dir=tmp_path / "outputs",
        sqlite_path=tmp_path / "provider_directory_truth_layer.sqlite3",
        as_of=datetime(2026, 6, 26),
    )

    workbook = load_workbook(tmp_path / "outputs/provider_directory_review_pack.xlsx")
    assert workbook.sheetnames == [
        "Summary",
        "Exceptions",
        "Duplicates",
        "Normalized Records",
        "Rules",
        "Review Notes",
    ]


def test_pipeline_writes_remediation_workbook_when_payer_listings_are_explicit(tmp_path):
    repo = Path(__file__).resolve().parents[1]
    run_pipeline(
        input_csv=repo / "data/raw/sample_provider_directory.csv",
        config_path=repo / "config/qa_rules.yml",
        output_dir=tmp_path / "outputs",
        sqlite_path=tmp_path / "provider_directory_truth_layer.sqlite3",
        as_of=datetime(2026, 6, 26),
        payer_listings_csv=repo / "data/raw/synthetic_payer_directory_listings.csv",
    )

    workbook = load_workbook(tmp_path / "outputs/provider_directory_review_pack.xlsx")
    assert workbook.sheetnames == [
        "Summary",
        "Exceptions",
        "Duplicates",
        "Normalized Records",
        "Rules",
        "Review Notes",
        "Payer Listings",
        "Listing Identity",
        "Field Precedence",
        "Reference Values",
        "Field Comparisons",
        "Discrepancies",
        "Remediation Risk",
    ]


def test_pipeline_writes_inspectable_remediation_outputs_and_sqlite_tables(tmp_path):
    repo = Path(__file__).resolve().parents[1]
    run_pipeline(
        input_csv=repo / "data/raw/sample_provider_directory.csv",
        config_path=repo / "config/qa_rules.yml",
        output_dir=tmp_path / "outputs",
        sqlite_path=tmp_path / "provider_directory_truth_layer.sqlite3",
        as_of=datetime(2026, 6, 26),
        payer_listings_csv=repo / "data/raw/synthetic_payer_directory_listings.csv",
    )

    expected_files = {
        "synthetic_payer_directory_listings.csv",
        "payer_listing_identity_results.csv",
        "field_source_precedence.csv",
        "provider_field_reference_values.csv",
        "payer_listing_field_comparisons.csv",
        "payer_directory_discrepancies.csv",
        "directory_remediation_risk_summary.csv",
    }
    assert {path.name for path in (tmp_path / "outputs").iterdir()} >= expected_files

    with sqlite3.connect(tmp_path / "provider_directory_truth_layer.sqlite3") as conn:
        tables = {
            row[0]
            for row in conn.execute("select name from sqlite_master where type = 'table'").fetchall()
        }
    assert {
        "synthetic_payer_directory_listings",
        "payer_listing_identity_results",
        "field_source_precedence",
        "provider_field_reference_values",
        "payer_listing_field_comparisons",
        "payer_directory_discrepancies",
        "directory_remediation_risk_summary",
    }.issubset(tables)


def test_custom_run_without_payer_listings_removes_remediation_artifacts(tmp_path):
    repo = Path(__file__).resolve().parents[1]
    output_dir = tmp_path / "outputs"
    sqlite_path = tmp_path / "provider_directory_truth_layer.sqlite3"
    common = {
        "input_csv": repo / "data/raw/sample_provider_directory.csv",
        "config_path": repo / "config/qa_rules.yml",
        "output_dir": output_dir,
        "sqlite_path": sqlite_path,
        "as_of": datetime(2026, 6, 26),
    }
    run_pipeline(
        **common,
        payer_listings_csv=repo / "data/raw/synthetic_payer_directory_listings.csv",
    )
    result = run_pipeline(**common)

    assert result.remediation is None
    assert not any((output_dir / filename).exists() for filename in REMEDIATION_OUTPUTS)
    assert "## Payer Directory Verification" not in (output_dir / "operations_review_memo.md").read_text()
    workbook = load_workbook(output_dir / "provider_directory_review_pack.xlsx")
    assert not {
        "Payer Listings",
        "Listing Identity",
        "Field Precedence",
        "Reference Values",
        "Field Comparisons",
        "Discrepancies",
        "Remediation Risk",
    }.intersection(workbook.sheetnames)
    with sqlite3.connect(sqlite_path) as conn:
        tables = {row[0] for row in conn.execute("select name from sqlite_master where type = 'table'")}
    assert not REMEDIATION_TABLES.intersection(tables)


def test_public_or_custom_input_does_not_enable_payer_remediation(tmp_path):
    input_csv = tmp_path / "public_npi_style_input.csv"
    pd.DataFrame(
        [
            {
                "record_id": "PUBLIC-00001",
                "npi": "1234567890",
                "provider_name": "Public Reference Provider",
                "specialty": "Primary Care",
                "address": "100 Public Way",
                "city": "Nashville",
                "state": "TN",
                "zip": "37201",
                "phone": "615-555-0100",
                "accepting_new_patients": "unknown",
                "source": "NPI Registry API",
                "status": "active",
                "last_verified_date": "2026-06-01",
            }
        ]
    ).to_csv(input_csv, index=False)
    repo = Path(__file__).resolve().parents[1]
    output_dir = tmp_path / "outputs"
    sqlite_path = tmp_path / "provider_directory_truth_layer.sqlite3"

    result = run_pipeline(
        input_csv=input_csv,
        config_path=repo / "config/qa_rules.yml",
        output_dir=output_dir,
        sqlite_path=sqlite_path,
        as_of=datetime(2026, 6, 26),
    )

    assert result.remediation is None
    assert not (output_dir / "payer_listing_field_comparisons.csv").exists()
    with sqlite3.connect(sqlite_path) as conn:
        tables = {row[0] for row in conn.execute("select name from sqlite_master where type = 'table'")}
    assert "payer_listing_field_comparisons" not in tables


def test_explicit_remediation_csv_evidence_is_byte_identical_across_complete_runs(tmp_path):
    repo = Path(__file__).resolve().parents[1]
    common = {
        "input_csv": repo / "data/raw/sample_provider_directory.csv",
        "config_path": repo / "config/qa_rules.yml",
        "as_of": datetime(2026, 6, 26),
        "payer_listings_csv": repo / "data/raw/synthetic_payer_directory_listings.csv",
    }
    first_output = tmp_path / "first"
    second_output = tmp_path / "second"
    run_pipeline(
        **common,
        output_dir=first_output,
        sqlite_path=tmp_path / "first.sqlite3",
    )
    run_pipeline(
        **common,
        output_dir=second_output,
        sqlite_path=tmp_path / "second.sqlite3",
    )

    evidence_files = {
        "synthetic_payer_directory_listings.csv",
        "payer_listing_identity_results.csv",
        "field_source_precedence.csv",
        "provider_field_reference_values.csv",
        "payer_listing_field_comparisons.csv",
        "payer_directory_discrepancies.csv",
        "directory_remediation_risk_summary.csv",
    }
    for filename in evidence_files:
        assert (first_output / filename).read_bytes() == (second_output / filename).read_bytes()
