"""Durable service contracts; every test owns its database and retained bytes."""
import csv
import io
import json
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook

from provider_directory_truth_layer import intake_store
from provider_directory_truth_layer.intake_store import Conflict, IntakeStore


@pytest.fixture
def store(tmp_path):
    return IntakeStore(tmp_path / 'directory.sqlite3', tmp_path / 'runtime')


def provider(**changes):
    return {'record_id': '001', 'npi': '1000000001', 'provider_name': 'Synthetic Clinic',
            'address': '10 Fiction St', 'city': 'Example', 'state': 'TN', 'zip': '01234',
            'phone': '5550100000', 'location_id': 'A', **changes}


def csv_bytes(rows):
    output = io.StringIO()
    fields = list(dict.fromkeys(key for row in rows for key in row))
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode()


def accept(store, rows, source='clinic:bounded-A'):
    receipt = store.receive(csv_bytes(rows), 'synthetic.csv', source=source)
    return store.commit(receipt['id'], {'actor': 'operator', 'reason': 'Reviewed preview', 'policy': 'valid_rows'})


def counts(store):
    with store.db() as conn:
        return {table: conn.execute(f'select count(*) from {table}').fetchone()[0]
                for table in ['artifact', 'receipt', 'attempt', 'observation', 'work_case', 'work_event']}


@pytest.mark.parametrize('filetype', ['csv', 'xlsx'])
def test_unknown_headers_require_mapping_and_partial_acceptance_preserves_quarantine(store, filetype):
    good = provider()
    bad = provider(record_id='002', npi='')
    headers = {key: 'Supplier ' + key for key in good}
    rows = [{headers[key]: value for key, value in row.items()} for row in [good, bad]]
    payload = csv_bytes(rows)
    if filetype == 'xlsx':
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Roster'
        sheet.append(list(rows[0]))
        for row in rows:
            sheet.append(list(row.values()))
        buffer = io.BytesIO()
        workbook.save(buffer)
        payload = buffer.getvalue()
    receipt = store.receive(payload, f'synthetic.{filetype}')
    assert receipt['state'] == 'needs_mapping'
    with pytest.raises(ValueError, match='mapped preview'):
        store.commit(receipt['id'], {'actor': 'op', 'reason': 'premature', 'policy': 'valid_rows'})
    receipt = store.mapping(receipt['id'], {'mapping': headers})
    assert receipt['preview']['valid_count'] == 1
    assert receipt['preview']['quarantined_count'] == 1
    assert 'Missing npi' in receipt['preview']['rows'][1]['problems']
    accepted = store.commit(receipt['id'], {'actor': 'op', 'reason': 'Accept valid only', 'policy': 'valid_rows'})
    assert len(accepted['observations']) == 1
    assert accepted['observations'][0]['normalized']['zip'] == '01234'
    assert accepted['observations'][0]['original']['Supplier record_id'] == '001'
    assert store.artifacts.joinpath(accepted['artifact_sha']).read_bytes() == payload
    assert IntakeStore(store.path, store.root).detail(receipt['id'])['preview']['quarantined_count'] == 1


def test_repeated_delivery_retains_receipt_reuses_artifact_and_interpretation(store):
    first = accept(store, [provider()])
    before = counts(store)
    repeated = accept(store, [provider()])
    assert repeated['state'] == 'duplicate_delivery'
    assert repeated['duplicate_of'] == first['id']
    assert repeated['artifact_sha'] == first['artifact_sha']
    assert repeated['observations'] == first['observations']
    after = counts(store)
    assert after['receipt'] == before['receipt'] + 1
    assert after['attempt'] == before['attempt'] + 1
    for table in ['artifact', 'observation', 'work_case']:
        assert after[table] == before[table]


def test_changed_mapping_creates_attempt_and_new_interpretation(store):
    receipt = accept(store, [provider(alternate_phone='5550100009')])
    mapping = receipt['preview']['mapping'].copy()
    mapping['phone'] = 'alternate_phone'
    revised = store.mapping(receipt['id'], {'mapping': mapping})
    assert len(revised['attempts']) == 2
    assert revised['attempts'][0]['fingerprint'] != revised['attempts'][1]['fingerprint']
    revised = store.commit(receipt['id'], {'actor': 'op', 'reason': 'Supplier mapping corrected', 'policy': 'valid_rows'})
    assert revised['observations'][0]['normalized']['phone'] == '5550100009'
    assert counts(store)['observation'] == 2


def test_reordered_changed_evidence_preserves_case_identity_owner_and_initial_evidence(store):
    first = accept(store, [provider(), provider(record_id='002', location_id='B', address='20 Fiction St')])
    first_case = next(case for case in first['linked_cases'] if case['location_id'].endswith(':A'))
    case = store.case(first_case['id'])
    original = case['evidence']
    store.event(case['id'], {'action': 'assign', 'owner': 'reviewer', 'actor': 'op', 'reason': 'Follow up', 'expected_version': 0})
    later = accept(store, [provider(record_id='002', location_id='B', address='20 Fiction St'), provider(phone='5550100008')])
    assert {case['id'] for case in later['linked_cases']} == {case['id'] for case in first['linked_cases']}
    refreshed = store.case(case['id'])
    assert refreshed['owner'] == 'reviewer'
    assert refreshed['evidence'] == original
    assert refreshed['evidence']['normalized']['phone'] == '5550100000'
    assert refreshed['later_evidence'][0]['normalized']['phone'] == '5550100008'
    assert refreshed['history'][0]['action'] == 'assign'


def test_source_scope_is_distinct_and_later_absence_does_not_terminate(store):
    first = accept(store, [provider(), provider(record_id='002', location_id='B')])
    accept(store, [provider()], source='clinic:bounded-B')
    assert counts(store)['work_case'] == 3
    accept(store, [provider(phone='5550100008')])
    for case in first['linked_cases']:
        assert store.case(case['id'])['status'] == 'new'
    assert counts(store)['observation'] == 4


def test_restart_recovers_processing_marker_without_committing_partial_evidence(store):
    receipt = store.receive(csv_bytes([provider()]), 'synthetic.csv')
    with store.db() as conn:
        conn.execute("update attempt set state='processing' where id=?", (receipt['attempt_id'],))
    restarted = IntakeStore(store.path, store.root)
    assert restarted.detail(receipt['id'])['state'] == 'interrupted'
    assert counts(restarted)['observation'] == 0
    assert restarted.mapping(receipt['id'], {})['state'] == 'preview_ready'


def test_commit_rolls_back_all_observations_and_cases_on_mid_batch_failure(store, monkeypatch):
    receipt = store.receive(csv_bytes([provider(), provider(record_id='002')]), 'synthetic.csv')
    original = store.link_work
    calls = []
    def fail_second(conn, source, observation):
        calls.append(observation)
        if len(calls) == 2:
            raise RuntimeError('Injected interrupted commit')
        return original(conn, source, observation)
    monkeypatch.setattr(store, 'link_work', fail_second)
    with pytest.raises(RuntimeError, match='Injected'):
        store.commit(receipt['id'], {'actor': 'op', 'reason': 'accept', 'policy': 'valid_rows'})
    assert counts(store)['observation'] == counts(store)['work_case'] == 0
    assert store.detail(receipt['id'])['state'] == 'preview_ready'
    monkeypatch.setattr(store, 'link_work', original)
    assert len(store.commit(receipt['id'], {'actor': 'op', 'reason': 'retry', 'policy': 'valid_rows'})['observations']) == 2


def test_expected_version_conflict_preserves_event_and_projection(store):
    case_id = accept(store, [provider()])['linked_cases'][0]['id']
    updated = store.event(case_id, {'actor': 'op', 'reason': 'Start', 'action': 'triage', 'expected_version': 0})
    with pytest.raises(Conflict, match='changed'):
        store.event(case_id, {'actor': 'other', 'reason': 'Stale draft', 'action': 'defer', 'expected_version': 0})
    assert store.case(case_id) == updated


def test_event_insert_failure_rolls_back_projection_and_history_is_append_only(store):
    case_id = accept(store, [provider()])['linked_cases'][0]['id']
    original = store.case(case_id)
    with store.db() as conn:
        conn.execute("create trigger injected_failure before insert on work_event begin select raise(abort,'Injected event failure'); end")
    with pytest.raises(sqlite3.IntegrityError, match='Injected'):
        store.event(case_id, {'actor': 'op', 'reason': 'Start', 'action': 'triage', 'expected_version': 0})
    assert store.case(case_id) == original
    with store.db() as conn:
        conn.execute('drop trigger injected_failure')
    store.event(case_id, {'actor': 'op', 'reason': 'Start', 'action': 'triage', 'expected_version': 0})
    with store.db() as conn:
        for statement in ['update work_event set reason=\'rewritten\'', 'delete from work_event']:
            with pytest.raises(sqlite3.IntegrityError, match='Append-only'):
                conn.execute(statement)


def test_same_provider_multiple_locations_never_compares_cross_site_phones(store):
    receipt = accept(store, [provider(), provider(record_id='002', location_id='B', phone='5550100008')])
    assert len(receipt['observations']) == len(receipt['linked_cases']) == 2
    assert all(case['kind'] == 'qa' for case in receipt['linked_cases'])
    assert len({row['location_id'] for row in receipt['observations']}) == 2


def test_location_relationship_requires_explicit_actor_reason_and_observed_npi(store):
    receipt = accept(store, [provider(location_id='')])
    observation = receipt['observations'][0]
    payload = {'observation_id': observation['id'], 'provider_id': 'NPI:1000000001',
               'location_id': 'confirmed:office-A', 'actor': 'op', 'reason': 'Confirmed source location'}
    with pytest.raises(ValueError, match='reason'):
        store.relationship({**payload, 'reason': ''})
    with pytest.raises(ValueError, match='exact observed NPI'):
        store.relationship({**payload, 'provider_id': 'NPI:1000000002'})
    store.relationship(payload)
    detail = store.detail(receipt['id'])
    assert detail['observations'][0]['location_id'] == 'confirmed:office-A'
    assert detail['observations'][0]['original'] == observation['original']
    with store.db() as conn:
        row = dict(conn.execute('select * from relationship').fetchone())
        assert row['actor'] == 'op' and row['reason'] == payload['reason']


def test_legacy_migration_preserves_tables_events_and_creates_original_backup(tmp_path):
    path = tmp_path / 'legacy.sqlite3'
    with sqlite3.connect(path) as conn:
        conn.executescript("create table review_event(id text primary key, reason text); insert into review_event values('legacy-event','Original reason');")
    store = IntakeStore(path, tmp_path / 'runtime')
    backup = path.with_name(path.name + '.pre-intake-v1.bak')
    assert backup.exists()
    with sqlite3.connect(backup) as conn:
        assert conn.execute('select * from review_event').fetchall() == [('legacy-event', 'Original reason')]
        assert conn.execute('pragma user_version').fetchone()[0] == 0
        assert conn.execute("select count(*) from sqlite_master where name='artifact'").fetchone()[0] == 0
    with store.db() as conn:
        assert tuple(conn.execute('select * from review_event').fetchone()) == ('legacy-event', 'Original reason')
        assert conn.execute('pragma user_version').fetchone()[0] == intake_store.VERSION


def test_future_schema_is_rejected_without_downgrading(tmp_path):
    path = tmp_path / 'future.sqlite3'
    with sqlite3.connect(path) as conn:
        conn.execute('pragma user_version=999')
    with pytest.raises(ValueError, match='future schema'):
        IntakeStore(path, tmp_path / 'runtime')
    with sqlite3.connect(path) as conn:
        assert conn.execute('pragma user_version').fetchone()[0] == 999


def test_failed_migration_rolls_back_schema_and_keeps_legacy_data(tmp_path, monkeypatch):
    path = tmp_path / 'legacy.sqlite3'
    with sqlite3.connect(path) as conn:
        conn.executescript("create table legacy(id integer); insert into legacy values(7);")
    monkeypatch.setattr(intake_store, 'SCHEMA', ['create table staged(id integer); invalid SQL;'])
    with pytest.raises(sqlite3.OperationalError):
        IntakeStore(path, tmp_path / 'runtime')
    with sqlite3.connect(path) as conn:
        assert conn.execute('select * from legacy').fetchall() == [(7,)]
        assert conn.execute('pragma user_version').fetchone()[0] == 0
        assert conn.execute("select count(*) from sqlite_master where name='staged'").fetchone()[0] == 0


@pytest.mark.parametrize('profile', ['nppes', 'cms', 'flat'])
def test_native_saved_json_profiles_commit_with_original_response_grain(store, profile):
    fixtures = Path(__file__).parent / 'fixtures' / 'intake'
    data = json.loads((fixtures / f'{profile}-synthetic.json').read_text())
    if profile == 'nppes':
        data['results'][0]['addresses'][0]['telephone_number'] = '5550100000'
    elif profile == 'flat':
        data[0].update(address='10 Fiction St', city='Example', state='TN', phone='5550100000')
    payload = json.dumps(data).encode()
    receipt = store.receive(payload, 'saved-synthetic.json', source=f'{profile}:bounded', profile=profile)
    assert receipt['state'] == 'preview_ready'
    assert receipt['preview']['valid_count'] >= 1
    committed = store.commit(receipt['id'], {'actor': 'op', 'reason': 'Reviewed saved response', 'policy': 'valid_rows'})
    assert len(committed['observations']) == receipt['preview']['valid_count']
    row = committed['observations'][0]
    assert row['human_verified_at'] is None
    assert row['native_id']
    if profile == 'nppes':
        assert len(row['normalized']['locations']) == 3
        assert row['original']['unknown_future_field'] == {'keep': True}
        assert row['source_updated_at'] == '2026-10-01'
    if profile == 'cms':
        assert row['native_id'] == '1000000001|synthetic-enroll-1|synthetic-group|synthetic-address-A'
    assert store.artifacts.joinpath(receipt['artifact_sha']).read_bytes() == payload


def test_sqlite_integrity_foreign_keys_and_artifact_tampering_detection(store):
    receipt = accept(store, [provider()])
    with store.db() as conn:
        assert conn.execute('pragma integrity_check').fetchone()[0] == 'ok'
        assert conn.execute('pragma foreign_key_check').fetchall() == []
        assert conn.execute('pragma foreign_keys').fetchone()[0] == 1
        with pytest.raises(sqlite3.IntegrityError):
            conn.execute("insert into work_evidence values('missing-case','missing-observation')")
    store.artifacts.joinpath(receipt['artifact_sha']).write_bytes(b'changed bytes')
    retried = store.mapping(receipt['id'], {})
    assert retried['state'] == 'failed'
    assert 'digest mismatch' in retried['preview']['error']


def test_correction_verification_requires_observed_evidence_and_retains_reopen_history(store):
    case_id = accept(store, [provider(reference_phone='5550100009')])['linked_cases'][0]['id']
    steps = [('triage', {}), ('approve', {'proposal': '5550100009'}),
             ('submit', {'submission_reference': 'SIMULATED ticket-1'}),
             ('verify', {'observed': '5550100000', 'verification_note': 'Synthetic failed check'}),
             ('triage', {}), ('approve', {'proposal': '5550100008'}),
             ('submit', {'submission_reference': 'SIMULATED ticket-2'}),
             ('verify', {'observed': '(555) 010-0008', 'verification_note': 'Synthetic observed match'})]
    states = []
    for version, (action, fields) in enumerate(steps):
        case = store.event(case_id, {'actor': 'op', 'reason': 'Synthetic workflow',
                                    'expected_version': version, 'action': action, **fields})
        states.append(case['status'])
    assert states == ['triaged', 'approved', 'pending_verification', 'reopened', 'triaged', 'approved', 'pending_verification', 'verified']
    assert len(case['history']) == 8
    assert case['history'][1]['payload']['proposal'] == '5550100009'
    assert case['proposal'] == '5550100008'
    assert case['evidence']['normalized']['phone'] == '5550100000'


def test_acceptance_version_phone_validation_and_corrected_later_evidence(store):
    receipt = store.receive(csv_bytes([provider(reference_phone='5550100009')]), 'synthetic.csv')
    stale_attempt = receipt['attempt_id']
    newer = store.mapping(receipt['id'], {'delimiter':','})
    from provider_directory_truth_layer.intake_store import Conflict
    with pytest.raises(Conflict, match='Interpretation changed'):
        store.commit(receipt['id'], {'actor':'op','reason':'Stale preview','policy':'valid_rows','expected_attempt':stale_attempt})
    store.commit(receipt['id'], {'actor':'op','reason':'Reviewed','policy':'valid_rows','expected_attempt':newer['attempt_id']})
    case = store.work()['items'][0]
    case = store.event(case['id'], {'action':'triage','actor':'op','reason':'Review','expected_version':0})
    with pytest.raises(ValueError, match='ten digits'):
        store.event(case['id'], {'action':'approve','proposal':'invalid phone','actor':'op','reason':'Invalid','expected_version':1})
    changed = accept(store,[provider(reference_phone='5550100009',phone='5550100009')],source='clinic')
    original = store.case(case['id'])
    assert original['status'] == 'triaged'
    assert len(original['later_evidence']) == 1
    assert len(store.work()['items']) == 1
    assert changed['preview']['comparison']['changed'] == 1


def test_real_bounds_formulas_and_formatted_numeric_zip(store):
    from provider_directory_truth_layer.intake_store import MAX_BYTES
    with pytest.raises(ValueError,match='10 MiB'):
        store.receive(b'x'*(MAX_BYTES+1),'oversize.csv')
    many=csv_bytes([provider(record_id=str(i)) for i in range(20001)])
    assert store.receive(many,'over-row-limit.csv')['state']=='failed'
    workbook=Workbook();sheet=workbook.active
    values=provider()
    sheet.append(list(values))
    sheet.append(list(values.values()))
    cell=sheet.cell(2,list(values).index('zip')+1)
    cell.value=1234;cell.number_format='00000'
    sheet.append(list(values.values()))
    sheet.cell(3,list(values).index('phone')+1).value='=1+1'
    stream=io.BytesIO();workbook.save(stream)
    receipt=store.receive(stream.getvalue(),'numeric-zip.xlsx')
    assert receipt['preview']['rows'][0]['normalized']['zip']=='01234'
    assert receipt['preview']['valid_count']==1
    assert any('Formula' in issue for issue in receipt['preview']['rows'][1]['problems'])


def test_multi_location_public_crosswalk_selects_coherent_practice_not_mailing(store):
    data=[provider(locations=[{'address':'Mailing Box','purpose':'MAILING','phone':'5550101111'},{'address':'10 Practice St','purpose':'LOCATION','phone':'5550102222'},{'address':'20 Practice St','purpose':'LOCATION','phone':'5550103333'}])]
    receipt=store.receive(json.dumps(data).encode(),'multi.json')
    accepted=store.commit(receipt['id'],{'actor':'op','reason':'Review native evidence','policy':'valid_rows'})
    observation=accepted['observations'][0]
    base={'observation_id':observation['id'],'provider_id':'NPI:1000000001','location_id':'clinic:chosen','actor':'op','reason':'Confirmed practice location'}
    with pytest.raises(ValueError,match='Choose a native location'):
        store.relationship(base)
    with pytest.raises(ValueError,match='Mailing'):
        store.relationship({**base,'location_index':0})
    store.relationship({**base,'location_index':2})
    relationships=store.detail(receipt['id'])['observations'][0]['relationships']
    assert relationships[0]['location_index']==2
    case=next(store.case(c['id']) for c in store.work()['items'] if c['location_id']=='clinic:chosen')
    assert case['evidence']['normalized']['address']=='20 Practice St'
    assert case['evidence']['normalized']['phone']=='5550103333'


def test_version_one_upgrade_keeps_data_and_backup(tmp_path):
    path=tmp_path/'v1.db'
    with sqlite3.connect(path) as conn:
        conn.executescript(intake_store.SCHEMA[0]+'PRAGMA user_version=1;')
        conn.execute("insert into artifact values ('existing',1,'existing')")
    upgraded=IntakeStore(path,tmp_path)
    with upgraded.db() as conn:
        assert conn.execute('pragma user_version').fetchone()[0]==2
        assert conn.execute('select sha from artifact').fetchone()[0]=='existing'
        assert 'location_index' in [r[1] for r in conn.execute('pragma table_info(relationship)')]
    with sqlite3.connect(path.with_name('v1.db.pre-intake-v2.bak')) as backup:
        assert backup.execute('pragma user_version').fetchone()[0]==1


def test_restart_reports_retained_orphan_artifacts_without_deleting(store):
    (store.artifacts/'orphan.staged').write_bytes(b'original interrupted bytes')
    (store.artifacts/'unregistered').write_bytes(b'finalized before ledger write')
    restarted=IntakeStore(store.path,store.root)
    report=json.loads((store.root/'artifact-recovery.json').read_text())
    assert 'orphan.staged' in report['staged']
    assert 'unregistered' in report['unregistered']
    assert (restarted.artifacts/'orphan.staged').exists()
    assert restarted.inbox()['items']==[]
