"""Read-only durable-work handoff, independent of legacy pipeline workbooks."""
from __future__ import annotations

import io
import json
import re
import sqlite3
from contextlib import closing
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .intake_store import now


def snapshot(path: Path, case_id: str | None = None) -> dict:
    """Pin one WAL read snapshot, materialize it, then release before XLSX rendering."""
    with closing(sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True)) as conn:
        conn.row_factory = sqlite3.Row
        conn.execute('pragma query_only=ON')
        conn.execute('begin')
        where, args = ('where id=?', (case_id,)) if case_id is not None else ('', ())
        cases = [dict(r) for r in conn.execute(f'select * from work_case {where} order by id', args)]
        if case_id is not None and not cases:
            raise KeyError('Work case not found')
        # The first SELECT establishes the database snapshot; timestamp is descriptive.
        result = {'snapshot_at': now(), 'case_id': case_id, 'cases': cases,
                  'evidence': [], 'decisions': [], 'relationships': []}
        receipt_ids = set()
        for case in cases:
            case['evidence'] = json.loads(case['evidence'])
            receipt_ids.add(case['evidence']['receipt_id'])
            result['decisions'].extend(dict(r) for r in conn.execute(
                'select * from work_event where case_id=? order by id', (case['id'],)))
            for row in conn.execute('''select o.*, a.receipt_id from observation o
                    join work_evidence w on w.observation_id=o.id join attempt a on a.id=o.attempt_id
                    where w.case_id=? order by o.retrieved_at,o.id''', (case['id'],)):
                observation = dict(row)
                receipt_ids.add(observation['receipt_id'])
                observation['case_id'] = case['id']
                observation['role'] = 'initial' if observation['id'] == case['evidence']['observation_id'] else 'later'
                observation['original'] = json.loads(observation['original'])
                observation['normalized'] = json.loads(observation['normalized'])
                result['evidence'].append(observation)
                # A native observation can support several locations. Only label
                # relationships as this case's when both scope identities match.
                result['relationships'].extend({**dict(r), 'case_id': case['id']} for r in conn.execute(
                    '''select * from relationship where observation_id=?
                       and provider_id=? and location_id=? order by id''',
                    (observation['id'], case['evidence'].get('provider_id'), case['location_id'])))
        # Include repeat deliveries of the same source/profile/artifact, even when
        # acceptance reused an earlier interpretation instead of creating observations.
        related = set(receipt_ids)
        for rid in receipt_ids:
            related.update(r[0] for r in conn.execute('''select r.id from receipt r join receipt origin
                on r.source=origin.source and r.profile=origin.profile and r.artifact_sha=origin.artifact_sha
                where origin.id=?''', (rid,)))
        result['receipts'], result['attempts'] = [], []
        for rid in sorted(related):
            result['receipts'].append(dict(conn.execute('select * from receipt where id=?', (rid,)).fetchone()))
            result['attempts'].extend(dict(r) for r in conn.execute(
                'select * from attempt where receipt_id=? order by created_at,id', (rid,)))
        conn.rollback()
        return result


def handoff_workbook(path: Path, case_id: str | None = None) -> bytes:
    data = snapshot(path, case_id)
    workbook = Workbook()
    workbook.remove(workbook.active)
    payloads = []

    def payload(value):
        key = f'P{len(payloads) + 1:07d}'
        # ASCII JSON avoids XML controls and splitting a non-BMP character in Excel.
        encoded = json.dumps(value, ensure_ascii=True, sort_keys=True, allow_nan=False)
        for part, start in enumerate(range(0, len(encoded), 30000), 1):
            payloads.append([key, str(part), encoded[start:start + 30000]])
        return key

    def text(value):
        if value is None:
            return None
        value = str(value)
        if value == '':
            return '""'
        displayed = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]', lambda match: f'\\u{ord(match[0]):04x}', value)
        # Escaping controls expands the display; check its size, not the raw text.
        if len(displayed.encode('utf-16-le')) // 2 > 30000:
            return 'payload:' + payload(value)
        return displayed

    def sheet(name, headers, rows):
        tab = workbook.create_sheet(name)
        # Explicit string cells prevent formulas, hyperlinks, numeric ID coercion,
        # or dates being reinterpreted by Excel. JSON payload chunks are also text.
        for row in [headers, *rows]:
            if tab.max_row >= 1_048_576:
                raise ValueError('Handoff exceeds XLSX row capacity; export a selected case')
            tab.append([text(value) for value in row])
            for cell in tab[tab.max_row]:
                if cell.value is not None:
                    cell.data_type = 's'
                cell.number_format = '@'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        tab.freeze_panes = 'A2'
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='204D58')
        for index, header in enumerate(headers, 1):
            tab.column_dimensions[get_column_letter(index)].width = min(48, max(18, len(header) + 3))
        return tab

    sheet('Read me', ['item', 'meaning'], [
        ['format_version', '1'], ['snapshot_at_utc', data['snapshot_at']],
        ['scope', 'Selected case: ' + case_id if case_id is not None else 'All durable work cases; UI filters do not apply'],
        ['case_count', str(len(data['cases']))],
        ['consistency', 'One read-only SQLite transaction; snapshot pinned by first case SELECT. Rendering begins after reads finish. No workflow writes.'],
        ['coverage', 'All linked initial/later observations, events and relationships; no 50-observation/200-case UI cap. Legacy QA/remediation excluded.'],
        ['shared_receipts', 'Related deliveries match source/profile/artifact, including repeats. Attempts contain complete receipt previews, possibly including rows for other cases.'],
        ['values', 'listed_value/requested_value are initial source assertions; approved_value is the active human proposal; observed_value is the most recently asserted observation, including failed verification.'],
        ['dates', 'Source update, retrieval, receipt and human verification dates are separate. A source update is not human verification. Verification actions/notes are in Decisions.'],
        ['nulls', 'Blank cell = null; empty string displayed as two quotes. Payloads JSON is authoritative for exact null/empty/type fidelity. Dates and all identifiers are text.'],
        ['payloads', 'Join json_chunk by payload_id in numeric part order, then parse JSON. References preserve complete originals, previews, metadata and snapshots without Excel cell truncation.'],
        ['safety', 'All populated cells are literal text, including = + - @ prefixes. XML control characters in scalar displays are escaped; JSON payloads retain exact values. No formulas or external links.'],
        ['boundary', 'Local portfolio system. Designation is retained per receipt. Synthetic examples, public reference and locally asserted decisions do not certify provider truth or payer updates.'],
    ])
    cases = []
    for case in data['cases']:
        evidence = case['evidence']
        norm = evidence.get('normalized', {})
        cases.append([case['id'], case['kind'], evidence.get('provider_id'), norm.get('npi'), case['provider_name'],
            case['location_id'], norm.get('record_id'), norm.get('listing_id'), evidence.get('source'), evidence.get('designation'),
            case['field_name'], case['issue'], case['risk_level'], case['risk_reason'], case['status'], case['owner'], case['follow_up'],
            case['version'], evidence.get('listing_value'), evidence.get('reference_value'), case['proposal'], case['observed'],
            case['created_at'], case['updated_at'], evidence.get('receipt_id'), evidence.get('observation_id'), evidence.get('locator'), payload(case)])
    sheet('Cases', ['case_id','kind','provider_id','npi','provider_name','location_id','native_record_id','listing_id','source','designation',
        'field_name','issue','risk_level','risk_reason','status','owner','follow_up','version','listed_value','requested_value',
        'approved_value','observed_value','created_at','updated_at','initial_receipt_id','initial_observation_id','initial_locator','case_payload_id'], cases)
    receipts = {r['id']: r for r in data['receipts']}
    case_evidence = {case['id']: case['evidence'] for case in data['cases']}
    # Relationship history is ordered by ID; a repeated assertion about this
    # observation/case uses its latest selected practice location for the display.
    scopes = {(r['case_id'], r['observation_id']): r for r in data['relationships']}
    observations = []
    for observation in data['evidence']:
        r = receipts[observation['receipt_id']]
        initial = case_evidence[observation['case_id']]
        # A native observation may contain multiple practice addresses. Display the
        # explicitly selected initial scope, while retaining the complete native
        # observation and every relationship separately.
        scoped = initial if observation['role'] == 'initial' else observation
        if observation['role'] == 'later' and (relationship := scopes.get((observation['case_id'], observation['id']))):
            norm = observation['normalized']
            locator = observation['locator']
            index = relationship['location_index']
            if index is not None:
                locations = norm.get('locations', [])
                if not isinstance(index, int) or not 0 <= index < len(locations):
                    raise ValueError('Retained location relationship has no matching native address')
                norm = {**norm, **locations[index]}
                locator += '.' + locations[index].get('source_path', f'locations[{index}]')
            scoped = {**observation, 'provider_id': relationship['provider_id'],
                      'location_id': relationship['location_id'], 'normalized': norm, 'locator': locator}
        norm = scoped['normalized']
        observations.append([observation['case_id'], observation['role'], observation['id'], observation['attempt_id'], r['id'], r['label'],
            r['source'], r['designation'], r['artifact_sha'], observation['native_id'], scoped.get('provider_id'), scoped.get('location_id'),
            scoped['locator'], norm.get('address'), norm.get('city'), norm.get('state'), norm.get('zip'), norm.get('phone'), norm.get('reference_phone'),
            observation['source_updated_at'], observation['retrieved_at'], observation['human_verified_at'], observation['provider_id'],
            observation['location_id'], observation['locator'], payload(observation)])
    sheet('Evidence', ['case_id','role','observation_id','attempt_id','receipt_id','receipt_label','source','designation','artifact_sha',
        'native_id','provider_id','location_id','row_locator','address','city','state','zip','listed_phone','requested_phone',
        'source_updated_at','retrieved_at','human_verified_at','native_observation_provider_id','native_observation_location_id',
        'native_observation_locator','observation_payload_id'], observations)
    sheet('Receipts', ['receipt_id','label','source','designation','received_at','artifact_sha','profile','duplicate_of','receipt_payload_id'],
        [[r['id'],r['label'],r['source'],r['designation'],r['received_at'],r['artifact_sha'],r['profile'],r['duplicate_of'],payload(r)] for r in data['receipts']])
    sheet('Attempts', ['attempt_id','receipt_id','state','created_at','committed_at','decision','attempt_payload_id'],
        [[a['id'],a['receipt_id'],a['state'],a['created_at'],a['committed_at'],a['decision'],payload(a)] for a in data['attempts']])
    decisions = []
    for event in data['decisions']:
        inputs = json.loads(event['payload'])
        decisions.append([event['id'],event['case_id'],event['action'],event['actor'],event['reason'],event['from_status'],event['to_status'],
            event['occurred_at'], inputs.get('owner'), inputs.get('follow_up'), inputs.get('proposal'), inputs.get('submission_reference'),
            inputs.get('observed'), inputs.get('verification_note'), payload({**event, 'payload': inputs})])
    sheet('Decisions', ['event_id','case_id','action','actor','reason','from_status','to_status','occurred_at','owner_input','follow_up_input',
        'proposal_input','simulated_submission_reference','observed_input','verification_note','event_payload_id'], decisions)
    sheet('Relationships', ['relationship_id','case_id','observation_id','provider_id','location_id','location_index','actor','reason','occurred_at'],
        [[r['id'],r['case_id'],r['observation_id'],r['provider_id'],r['location_id'],r['location_index'],r['actor'],r['reason'],r['occurred_at']] for r in data['relationships']])
    sheet('Payloads', ['payload_id','part','json_chunk'], payloads)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
