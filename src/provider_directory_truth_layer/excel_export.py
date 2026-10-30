from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import RuleConfig
from .qa import QAResult
from .remediation import RemediationResult


def write_review_workbook(
    result: QAResult,
    config: RuleConfig,
    path: Path,
    remediation: RemediationResult | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    rule_rows = [
        {"rule": "required_fields", "value": ", ".join(config.required_fields)},
        {"rule": "phone_min_digits", "value": config.phone_min_digits},
        {"rule": "active_status_values", "value": ", ".join(sorted(config.active_status_values))},
        {"rule": "stale_days_threshold", "value": config.stale_days_threshold},
        {"rule": "duplicate_match_fields", "value": ", ".join(config.duplicate_match_fields)},
    ]
    if remediation is not None:
        rule_rows.append(
            {
                "rule": "remediation_comparison_fields",
                "value": ", ".join(config.remediation.comparison_fields),
            }
        )
    rules = pd.DataFrame(rule_rows)
    review_notes = pd.DataFrame(
        [
            {
                "note": "Review high priority exceptions first. Do not merge duplicate candidates without human confirmation."
            },
            {
                "note": "This workbook is a review aid, not a certification that provider data is correct."
            },
        ]
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        result.summary.to_excel(writer, sheet_name="Summary", index=False)
        result.exceptions.to_excel(writer, sheet_name="Exceptions", index=False)
        result.duplicate_candidates.to_excel(writer, sheet_name="Duplicates", index=False)
        result.clean_records.to_excel(writer, sheet_name="Normalized Records", index=False)
        rules.to_excel(writer, sheet_name="Rules", index=False)
        review_notes.to_excel(writer, sheet_name="Review Notes", index=False)
        if remediation is not None:
            remediation.payer_listings.to_excel(writer, sheet_name="Payer Listings", index=False)
            remediation.listing_identity_results.to_excel(
                writer, sheet_name="Listing Identity", index=False
            )
            remediation.field_source_precedence.to_excel(
                writer, sheet_name="Field Precedence", index=False
            )
            remediation.reference_values.to_excel(
                writer, sheet_name="Reference Values", index=False
            )
            remediation.field_comparisons.to_excel(
                writer, sheet_name="Field Comparisons", index=False
            )
            remediation.discrepancies.to_excel(writer, sheet_name="Discrepancies", index=False)
            remediation.risk_summary.to_excel(writer, sheet_name="Remediation Risk", index=False)

        for sheet in writer.book.worksheets:
            _format_sheet(sheet)


def _format_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        width = min(max(length + 2, 12), 48)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
